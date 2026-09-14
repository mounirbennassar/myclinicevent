"""Attendee management for event teams: list, walk-ins, edits, manual attendance, export, certificates."""

import io
from datetime import UTC, timedelta

from fastapi import APIRouter, BackgroundTasks, Depends, File, Query, Request, Response, UploadFile
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from .. import services as svc
from ..audit import audit
from ..db import get_db
from ..emailer import email_enabled, queue_certificate_email, queue_registration_email
from ..errors import ApiError
from ..importer import parse_attendees, profession_key
from ..models import Event, Registration, Scan
from ..schemas import ManualScanIn, RegistrationIn, RegistrationUpdate, ScanUpdate
from ..security import EventAccess, event_access
from ..validators import normalize_digits
from .scans import auto_issue_after_checkout

router = APIRouter(prefix="/events/{event_id}", tags=["registrations"])

ATTENDANCE_STATUSES = ("not_arrived", "inside", "checked_out", "no_checkout")


def _get_registration(db: Session, event: Event, registration_id: int) -> Registration:
    reg = db.get(Registration, registration_id)
    if reg is None or reg.event_id != event.id:
        raise ApiError(404, "registration_not_found", "Attendee not found.")
    return reg


def _detail(db: Session, access: EventAccess, reg: Registration) -> dict:
    event = access.event
    db.refresh(reg, ["scans"])
    summary = svc.summarize(event, [(s.direction, s.scanned_at) for s in reg.scans if not s.voided], reg)
    out = svc.registration_out(reg, summary, full_pii=access.is_manager)
    out["scans"] = [svc.scan_out(s) for s in reg.scans]
    out["email_sent_at"] = reg.email_sent_at
    out["certificate_blocker"] = svc.certificate_blocker(event, summary)
    if access.is_manager:
        out["pass_url"] = svc.pass_url(reg)
    return out


@router.get("/registrations")
def list_registrations(
    q: str = "",
    status: str = "all",
    source: str = "",
    sort: str = "recent",
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=500),
    access: EventAccess = Depends(event_access("scanner")),
    db: Session = Depends(get_db),
):
    rows = svc.event_attendance(db, access.event, include_cancelled=True)
    active = [(r, a) for r, a in rows if r.status != "cancelled"]
    counts = {
        "all": len(active),
        "cancelled": len(rows) - len(active),
        "eligible": sum(1 for _, a in active if a.eligible),
        "not_eligible": sum(1 for _, a in active if a.status != "not_arrived" and not a.eligible),
        **{s: sum(1 for _, a in active if a.status == s) for s in ATTENDANCE_STATUSES},
    }

    if status == "cancelled":
        pool = [x for x in rows if x[0].status == "cancelled"]
    elif status == "eligible":
        pool = [x for x in active if x[1].eligible]
    elif status == "not_eligible":
        pool = [x for x in active if x[1].status != "not_arrived" and not x[1].eligible]
    elif status in ATTENDANCE_STATUSES:
        pool = [x for x in active if x[1].status == status]
    else:
        pool = active

    if source in ("online", "walkin", "import"):
        pool = [x for x in pool if x[0].source == source]

    needle = normalize_digits(q.strip().lower())
    if needle:
        def matches(r: Registration) -> bool:
            fields = [r.full_name.lower(), r.email, r.ticket_code.lower(), r.mobile, r.scfhs_number.lower()]
            if access.is_manager:
                fields.append(r.national_id)
            return any(needle in f for f in fields)

        pool = [x for x in pool if matches(x[0])]

    if sort == "name":
        pool.sort(key=lambda x: x[0].full_name.lower())
    elif sort == "percent":
        pool.sort(key=lambda x: x[1].percent, reverse=True)
    elif sort == "arrival":
        pool.sort(key=lambda x: (x[1].first_in is None, x[1].first_in or x[0].created_at))
    else:
        pool.sort(key=lambda x: x[0].created_at, reverse=True)

    start = (page - 1) * page_size
    return {
        "items": [svc.registration_out(r, a, full_pii=access.is_manager) for r, a in pool[start:start + page_size]],
        "total": len(pool),
        "page": page,
        "page_size": page_size,
        "counts": counts,
    }


@router.post("/registrations", status_code=201)
def create_walkin(
    body: RegistrationIn,
    request: Request,
    background: BackgroundTasks,
    access: EventAccess = Depends(event_access("scanner")),
    db: Session = Depends(get_db),
):
    reg = svc.create_registration(db, access.event, body, source="walkin", actor=access.user)
    audit(db, actor=access.user, action="registration.walkin_created", entity_type="registration",
          entity_id=reg.id, event_id=access.event.id, request=request)
    db.commit()
    if queue_registration_email(background, reg, access.event):
        reg.email_sent_at = svc.utcnow()
        db.commit()
    out = _detail(db, access, reg)
    # The desk just created this person, so give them the link to share even if they're a scanner.
    out["pass_url"] = svc.pass_url(reg)
    return out


@router.post("/registrations/import")
async def import_registrations(
    request: Request,
    background: BackgroundTasks,
    file: UploadFile = File(...),
    notify: bool = False,
    access: EventAccess = Depends(event_access("manager")),
    db: Session = Depends(get_db),
):
    """Bulk pre-registration from a spreadsheet. Rows that fail validation or duplicate an existing
    attendee are skipped and reported; the rest are created (and optionally emailed their pass)."""
    data = await file.read()
    if len(data) > 5 * 1024 * 1024:
        raise ApiError(413, "file_too_large", "The file is larger than 5 MB.")
    rows = parse_attendees(file.filename or "", data)
    if not rows:
        raise ApiError(422, "empty_file", "No attendee rows were found in the file.")

    event = access.event
    created: list[Registration] = []
    skipped: list[dict] = []
    for row_number, raw in rows:
        profession = raw.get("profession", "")
        payload = RegistrationIn(
            full_name=raw.get("full_name", "")[:300],
            email=raw.get("email", "")[:300],
            mobile=raw.get("mobile", "")[:40],
            scfhs_number=raw.get("scfhs_number", "")[:60],
            national_id=raw.get("national_id", "")[:40],
            profession=profession_key(profession),
            consent=False,
        )
        try:
            created.append(svc.create_registration(db, event, payload, source="import", actor=access.user))
        except ApiError as err:
            detail = err.detail if isinstance(err.detail, dict) else {}
            skipped.append({
                "row": row_number,
                "name": raw.get("full_name", "")[:100],
                "code": detail.get("code", "error"),
                "fields": detail.get("fields") or {},
            })
    audit(db, actor=access.user, action="registrations.imported", entity_type="event", entity_id=event.id,
          event_id=event.id, details={"created": len(created), "skipped": len(skipped), "file": file.filename},
          request=request)
    db.commit()
    emailed = False
    if notify and created:
        for reg in created:
            emailed = queue_registration_email(background, reg, event)
        if emailed:
            for reg in created:
                reg.email_sent_at = svc.utcnow()
            db.commit()
    return {"created": len(created), "skipped": skipped, "total_rows": len(rows), "emailed": emailed}


_EXPORT_HEADERS = [
    "Ticket", "Full name", "Email", "Mobile", "SCFHS No.", "National ID / Iqama", "Profession", "Source",
    "Registration status", "Registered at", "Attendance status", "First check-in", "Last scan",
    "Attended (min)", "Required (min)", "Attendance %", "CME eligible", "Eligibility override", "Certificate code",
]


def _write_row(ws, values: list) -> None:
    ws.append(values)
    for cell in ws[ws.max_row]:
        # Store anything that looks like a formula as plain text.
        if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
            cell.data_type = "s"


@router.get("/registrations/export.xlsx")
def export_registrations(request: Request, access: EventAccess = Depends(event_access("manager")), db: Session = Depends(get_db)):
    event = access.event
    tz = svc.tz_of(event)
    rows = svc.event_attendance(db, event, include_cancelled=True)

    def fmt(dt):
        return dt.astimezone(tz).strftime("%Y-%m-%d %H:%M") if dt else ""

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendees"
    ws.append(_EXPORT_HEADERS)
    for r, a in rows:
        _write_row(ws, [
            r.ticket_code, r.full_name, r.email, r.mobile, r.scfhs_number, r.national_id, r.profession or "",
            r.source, r.status, fmt(r.created_at), a.status, fmt(a.first_in), fmt(a.last_scan_at),
            a.attended_seconds // 60, a.required_seconds // 60, a.percent, "Yes" if a.eligible else "No",
            "" if r.eligibility_override is None else ("Eligible" if r.eligibility_override else "Not eligible"),
            r.certificate_code or "",
        ])

    log = wb.create_sheet("Scan log")
    log.append(["Ticket", "Full name", "Direction", "Time", "Method", "Scanned by", "Voided", "Note"])
    scans = db.scalars(select(Scan).where(Scan.event_id == event.id).order_by(Scan.scanned_at)).all()
    names = {r.id: (r.ticket_code, r.full_name) for r, _ in rows}
    for s in scans:
        ticket, name = names.get(s.registration_id, ("", ""))
        _write_row(log, [ticket, name, s.direction, fmt(s.scanned_at), s.method,
                         s.scanned_by.full_name if s.scanned_by else "", "Yes" if s.voided else "", s.note or ""])

    header_fill = PatternFill("solid", fgColor="003868")
    for sheet in (ws, log):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            width = max(len(str(c.value or "")) for c in column)
            sheet.column_dimensions[column[0].column_letter].width = min(42, max(10, width + 2))

    buf = io.BytesIO()
    wb.save(buf)
    audit(db, actor=access.user, action="registrations.exported", entity_type="event", entity_id=event.id,
          event_id=event.id, details={"rows": len(rows)}, request=request)
    db.commit()
    filename = f"{event.slug}-attendees-{svc.utcnow().astimezone(tz):%Y%m%d-%H%M}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/registrations/{registration_id}")
def get_registration(
    registration_id: int, access: EventAccess = Depends(event_access("scanner")), db: Session = Depends(get_db)
):
    return _detail(db, access, _get_registration(db, access.event, registration_id))


@router.patch("/registrations/{registration_id}")
def update_registration(
    registration_id: int,
    body: RegistrationUpdate,
    request: Request,
    access: EventAccess = Depends(event_access("manager")),
    db: Session = Depends(get_db),
):
    reg = _get_registration(db, access.event, registration_id)
    data = body.model_dump(exclude_unset=True)
    identity = {k: v for k, v in data.items() if k in svc.REGISTRATION_FIELDS and v not in (None, "")}
    clean, errors = svc.clean_registration_fields(identity)
    if errors:
        raise ApiError(422, "validation", "Please check the highlighted fields.", errors)
    if "email" in clean or "national_id" in clean:
        svc.ensure_unique_identity(db, reg.event_id, clean.get("email"), clean.get("national_id"), exclude_id=reg.id)
    for key, value in clean.items():
        setattr(reg, key, value)
    if "profession" in data and not data["profession"]:
        reg.profession = None
    if data.get("status"):
        reg.status = data["status"]
    if "notes" in data:
        reg.notes = data["notes"] or None
    if "eligibility_override" in data:
        reg.eligibility_override = data["eligibility_override"]
    # Log which fields changed, not their values (they include ID numbers).
    details = {"fields": sorted(data)}
    if "eligibility_override" in data:
        details["eligibility_override"] = data["eligibility_override"]
    if "status" in data:
        details["status"] = data["status"]
    audit(db, actor=access.user, action="registration.updated", entity_type="registration", entity_id=reg.id,
          event_id=access.event.id, details=details, request=request)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise ApiError(409, "duplicate", "Another attendee already uses that email or ID number.") from None
    return _detail(db, access, reg)


@router.delete("/registrations/{registration_id}", status_code=204)
def delete_registration(
    registration_id: int,
    request: Request,
    access: EventAccess = Depends(event_access("manager")),
    db: Session = Depends(get_db),
):
    reg = _get_registration(db, access.event, registration_id)
    audit(db, actor=access.user, action="registration.deleted", entity_type="registration", entity_id=reg.id,
          event_id=access.event.id, details={"ticket_code": reg.ticket_code, "full_name": reg.full_name}, request=request)
    db.delete(reg)
    db.commit()
    return Response(status_code=204)


@router.post("/registrations/{registration_id}/scans", status_code=201)
def add_manual_scan(
    registration_id: int,
    body: ManualScanIn,
    request: Request,
    background: BackgroundTasks,
    access: EventAccess = Depends(event_access("manager")),
    db: Session = Depends(get_db),
):
    """Record a check-in/out by hand, e.g. a forgotten check-out. Bypasses the scanner's safety checks."""
    reg = _get_registration(db, access.event, registration_id)
    now = svc.utcnow()
    at = body.at or now
    if at.tzinfo is None:
        at = at.replace(tzinfo=svc.tz_of(access.event))
    at = at.astimezone(UTC)
    if at > now + timedelta(minutes=5):
        raise ApiError(422, "validation", "The time can't be in the future.", {"at": "future_time"})
    scan = Scan(event_id=access.event.id, registration_id=reg.id, direction=body.direction, scanned_at=at,
                method="admin", scanned_by_id=access.user.id, note=body.note)
    db.add(scan)
    db.flush()
    audit(db, actor=access.user, action="scan.manual_added", entity_type="scan", entity_id=scan.id,
          event_id=access.event.id, details={"direction": body.direction, "registration_id": reg.id}, request=request)
    if body.direction == "out":
        # A corrected check-out counts like a scanned one: hand out the certificate if it's now earned.
        auto_issue_after_checkout(
            db, background, request, access.event, reg, svc.registration_summary(db, access.event, reg),
            actor_id=access.user.id, via="manual_checkout",
        )
    db.commit()
    return _detail(db, access, reg)


@router.patch("/scans/{scan_id}")
def update_scan(
    scan_id: int,
    body: ScanUpdate,
    request: Request,
    access: EventAccess = Depends(event_access("manager")),
    db: Session = Depends(get_db),
):
    scan = db.get(Scan, scan_id)
    if scan is None or scan.event_id != access.event.id:
        raise ApiError(404, "scan_not_found", "Scan not found.")
    scan.voided = body.voided
    if body.note is not None:
        scan.note = body.note or None
    audit(db, actor=access.user, action="scan.voided" if body.voided else "scan.restored", entity_type="scan",
          entity_id=scan.id, event_id=access.event.id, request=request)
    db.commit()
    return _detail(db, access, scan.registration)


@router.post("/registrations/{registration_id}/resend")
def resend_pass(
    registration_id: int,
    request: Request,
    background: BackgroundTasks,
    access: EventAccess = Depends(event_access("manager")),
    db: Session = Depends(get_db),
):
    reg = _get_registration(db, access.event, registration_id)
    sent = queue_registration_email(background, reg, access.event)
    if sent:
        reg.email_sent_at = svc.utcnow()
    audit(db, actor=access.user, action="registration.pass_resent", entity_type="registration", entity_id=reg.id,
          event_id=access.event.id, request=request)
    db.commit()
    return {"email_enabled": sent}


@router.post("/registrations/{registration_id}/certificate")
def issue_certificate(
    registration_id: int,
    request: Request,
    background: BackgroundTasks,
    notify: bool = False,
    access: EventAccess = Depends(event_access("manager")),
    db: Session = Depends(get_db),
):
    event = access.event
    reg = _get_registration(db, event, registration_id)
    if not event.certificates_enabled:
        raise ApiError(422, "certificates_disabled", svc.CERTIFICATE_MESSAGES["certificates_disabled"])
    if not svc.registration_summary(db, event, reg).eligible:
        raise ApiError(422, "not_eligible", "This attendee hasn't met the attendance requirement. "
                                            "Set an eligibility override first to issue anyway.")
    if svc.issue_certificate(db, reg):
        audit(db, actor=access.user, action="certificate.issued", entity_type="registration", entity_id=reg.id,
              event_id=event.id, details={"via": "manager"}, request=request)
    db.commit()
    if notify:
        queue_certificate_email(background, reg, event)
    return _detail(db, access, reg)


@router.delete("/registrations/{registration_id}/certificate")
def revoke_certificate(
    registration_id: int,
    request: Request,
    access: EventAccess = Depends(event_access("manager")),
    db: Session = Depends(get_db),
):
    """Withdraw a certificate (e.g. after voiding a scan). Its code stops verifying immediately."""
    reg = _get_registration(db, access.event, registration_id)
    if reg.certificate_code:
        audit(db, actor=access.user, action="certificate.revoked", entity_type="registration", entity_id=reg.id,
              event_id=access.event.id, details={"code": reg.certificate_code}, request=request)
        svc.revoke_certificate(reg)
        db.commit()
    return _detail(db, access, reg)


@router.post("/certificates/issue-all")
def issue_all_certificates(
    request: Request,
    background: BackgroundTasks,
    notify: bool = True,
    access: EventAccess = Depends(event_access("manager")),
    db: Session = Depends(get_db),
):
    event = access.event
    if not event.certificates_enabled:
        raise ApiError(422, "certificates_disabled", svc.CERTIFICATE_MESSAGES["certificates_disabled"])
    issued = []
    for reg, summary in svc.event_attendance(db, event):
        if summary.eligible and svc.issue_certificate(db, reg):
            issued.append(reg)
    audit(db, actor=access.user, action="certificate.bulk_issued", entity_type="event", entity_id=event.id,
          event_id=event.id, details={"count": len(issued)}, request=request)
    db.commit()
    if notify:
        for reg in issued:
            queue_certificate_email(background, reg, event)
    return {"issued": len(issued), "email_enabled": email_enabled()}

"""The sponsor's own portal: profile, leads, lead scanner, team badges."""

import io
import re

from fastapi import APIRouter, BackgroundTasks, Depends, Query, Request, Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from .. import services as svc
from ..audit import audit
from ..db import get_db
from ..emailer import queue_sponsor_member_email
from ..errors import ApiError
from ..models import BadgeScan, Registration, Sponsor, SponsorLead, SponsorMember, User
from ..schemas import LeadScanIn, LeadUpdate, SponsorMemberIn, SponsorProfileUpdate
from ..security import require_roles
from ..sponsors import add_member, ensure_portal_user, sponsor_stats, upsert_lead
from ..validators import clean_mobile, normalize_digits

router = APIRouter(prefix="/sponsor", tags=["sponsor portal"])


def current_sponsor(user: User = Depends(require_roles("sponsor")), db: Session = Depends(get_db)) -> Sponsor:
    sponsor = db.get(Sponsor, user.sponsor_id) if user.sponsor_id else None
    if sponsor is None:
        raise ApiError(403, "sponsor_not_found", "Your account isn't linked to a sponsor.")
    if sponsor.status != "approved":
        raise ApiError(403, "sponsor_not_approved", "Your sponsorship hasn't been approved yet.")
    return sponsor


@router.get("/me")
def me(sponsor: Sponsor = Depends(current_sponsor), db: Session = Depends(get_db)):
    return {
        "sponsor": svc.sponsor_out(sponsor),
        "event": svc.public_event_out(sponsor.event),
        "stats": sponsor_stats(db, sponsor),
        "booth_url": svc.booth_url(sponsor),
    }


@router.patch("/profile")
def update_profile(
    body: SponsorProfileUpdate,
    request: Request,
    user: User = Depends(require_roles("sponsor")),
    sponsor: Sponsor = Depends(current_sponsor),
    db: Session = Depends(get_db),
):
    data = body.model_dump(exclude_unset=True)
    if "contact_mobile" in data and data["contact_mobile"]:
        mobile, err = clean_mobile(data["contact_mobile"])
        if err:
            raise ApiError(422, "validation", "Enter a valid mobile number.", {"contact_mobile": err})
        data["contact_mobile"] = mobile
    for key, value in data.items():
        if key == "contact_mobile" and not value:
            continue
        setattr(sponsor, key, (value or "").strip() or None if isinstance(value, str) else value)
    audit(db, actor=user, action="sponsor.profile_updated", entity_type="sponsor", entity_id=sponsor.id,
          event_id=sponsor.event_id, details={"fields": sorted(data)}, request=request)
    db.commit()
    return svc.sponsor_out(sponsor)


# ---------- leads


def _leads_query(sponsor: Sponsor):
    return select(SponsorLead).options(selectinload(SponsorLead.registration)).where(SponsorLead.sponsor_id == sponsor.id)


@router.get("/leads")
def list_leads(
    q: str = "",
    method: str = "",
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=500),
    sponsor: Sponsor = Depends(current_sponsor),
    db: Session = Depends(get_db),
):
    leads = db.scalars(_leads_query(sponsor).order_by(SponsorLead.captured_at.desc())).all()
    counts = {
        "all": len(leads),
        "booth_qr": sum(1 for lead in leads if lead.method == "booth_qr"),
        "badge_scan": sum(1 for lead in leads if lead.method == "badge_scan"),
        "with_contact": sum(1 for lead in leads if lead.consent or lead.registration.sponsor_consent),
    }
    if method in ("booth_qr", "badge_scan"):
        leads = [lead for lead in leads if lead.method == method]
    elif method == "with_contact":
        leads = [lead for lead in leads if lead.consent or lead.registration.sponsor_consent]
    needle = q.strip().lower()
    if needle:
        leads = [
            lead for lead in leads
            if needle in lead.registration.full_name.lower()
            or needle in (lead.note or "").lower()
            or needle in (lead.registration.profession or "")
        ]
    start = (page - 1) * page_size
    return {"items": [svc.lead_out(lead) for lead in leads[start:start + page_size]], "total": len(leads),
            "page": page, "page_size": page_size, "counts": counts}


@router.patch("/leads/{lead_id}")
def update_lead(
    lead_id: int, body: LeadUpdate, sponsor: Sponsor = Depends(current_sponsor), db: Session = Depends(get_db)
):
    lead = db.get(SponsorLead, lead_id)
    if lead is None or lead.sponsor_id != sponsor.id:
        raise ApiError(404, "lead_not_found", "Lead not found.")
    data = body.model_dump(exclude_unset=True)
    if "rating" in data:
        lead.rating = data["rating"]
    if "note" in data:
        lead.note = (data["note"] or "").strip() or None
    db.commit()
    return svc.lead_out(lead)


@router.post("/leads/scan")
def scan_lead(
    body: LeadScanIn,
    request: Request,
    user: User = Depends(require_roles("sponsor")),
    sponsor: Sponsor = Depends(current_sponsor),
    db: Session = Depends(get_db),
):
    """A booth rep scans an attendee's pass (or types their ticket code)."""
    raw = body.code.strip()
    if raw.upper().startswith(svc.QR_PREFIX):
        reg = db.scalar(select(Registration).where(Registration.qr_token == raw[len(svc.QR_PREFIX):].strip()))
    else:
        compact = re.sub(r"\s", "", normalize_digits(raw)).upper()
        ticket = compact if compact.startswith("MC-") else f"MC-{compact.removeprefix('MC')}"
        reg = db.scalar(select(Registration).where(Registration.ticket_code == ticket))
    if reg is None or reg.status == "cancelled":
        return {"result": "invalid", "lead": None}
    if reg.event_id != sponsor.event_id:
        return {"result": "wrong_event", "lead": None}
    lead, created = upsert_lead(db, sponsor, reg, method="badge_scan", consent=reg.sponsor_consent, captured_by=user)
    if created:
        audit(db, actor=user, action="sponsor.lead_scanned", entity_type="registration", entity_id=reg.id,
              event_id=sponsor.event_id, details={"sponsor_id": sponsor.id}, request=request)
    db.commit()
    db.refresh(lead)
    return {"result": "captured" if created else "already", "lead": svc.lead_out(lead)}


@router.get("/leads/export.xlsx")
def export_leads(sponsor: Sponsor = Depends(current_sponsor), db: Session = Depends(get_db)):
    tz = svc.tz_of(sponsor.event)
    leads = db.scalars(_leads_query(sponsor).order_by(SponsorLead.captured_at)).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append(["Captured at", "How", "Full name", "Profession", "Email", "Mobile", "Rating", "Note", "Contact shared"])
    for lead in leads:
        out = svc.lead_out(lead)
        a = out["attendee"]
        ws.append([
            lead.captured_at.astimezone(tz).strftime("%Y-%m-%d %H:%M"),
            "Booth visit" if lead.method == "booth_qr" else "Badge scan",
            a["full_name"], a["profession"] or "", a["email"] or "", a["mobile"] or "",
            lead.rating or "", lead.note or "", "Yes" if out["consent"] else "No",
        ])
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                cell.data_type = "s"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="003868")
    ws.freeze_panes = "A2"
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = min(42, max(12, max(len(str(c.value or "")) for c in column) + 2))
    buf = io.BytesIO()
    wb.save(buf)
    filename = f"{svc.slugify(sponsor.company_name) or 'sponsor'}-leads.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---------- team


@router.get("/members")
def list_members(sponsor: Sponsor = Depends(current_sponsor)):
    return [svc.member_out(m) for m in sponsor.members]


@router.post("/members", status_code=201)
def create_member(
    body: SponsorMemberIn,
    request: Request,
    background: BackgroundTasks,
    user: User = Depends(require_roles("sponsor")),
    sponsor: Sponsor = Depends(current_sponsor),
    db: Session = Depends(get_db),
):
    member, temporary, emailed = add_member(db, sponsor, body, background=background)
    audit(db, actor=user, action="sponsor.member_added", entity_type="sponsor_member", entity_id=member.id,
          event_id=sponsor.event_id, details={"portal_access": body.portal_access}, request=request)
    db.commit()
    return {"member": svc.member_out(member), "temporary_password": temporary, "emailed": emailed}


@router.post("/members/{member_id}/resend")
def resend_badge(
    member_id: int,
    background: BackgroundTasks,
    sponsor: Sponsor = Depends(current_sponsor),
):
    member = next((m for m in sponsor.members if m.id == member_id), None)
    if member is None:
        raise ApiError(404, "member_not_found", "Team member not found.")
    return {"emailed": queue_sponsor_member_email(background, member, sponsor, sponsor.event, None)}


@router.post("/members/{member_id}/portal-access")
def grant_portal_access(
    member_id: int,
    request: Request,
    background: BackgroundTasks,
    user: User = Depends(require_roles("sponsor")),
    sponsor: Sponsor = Depends(current_sponsor),
    db: Session = Depends(get_db),
):
    member = next((m for m in sponsor.members if m.id == member_id), None)
    if member is None:
        raise ApiError(404, "member_not_found", "Team member not found.")
    _, temporary = ensure_portal_user(db, sponsor, member)
    audit(db, actor=user, action="sponsor.portal_access_granted", entity_type="sponsor_member", entity_id=member.id,
          event_id=sponsor.event_id, request=request)
    db.commit()
    emailed = queue_sponsor_member_email(background, member, sponsor, sponsor.event, temporary) if temporary else False
    return {"member": svc.member_out(member), "temporary_password": temporary, "emailed": emailed}


@router.delete("/members/{member_id}", status_code=204)
def remove_member(
    member_id: int,
    request: Request,
    user: User = Depends(require_roles("sponsor")),
    sponsor: Sponsor = Depends(current_sponsor),
    db: Session = Depends(get_db),
):
    member = next((m for m in sponsor.members if m.id == member_id), None)
    if member is None:
        raise ApiError(404, "member_not_found", "Team member not found.")
    if member.is_contact:
        raise ApiError(422, "contact_member", "The main contact can't be removed.")
    if member.user_id == user.id:
        raise ApiError(422, "self_change", "You can't remove yourself.")
    if member.user_id:
        portal_user = db.get(User, member.user_id)
        if portal_user is not None:
            portal_user.is_active = False
            portal_user.token_version += 1
    audit(db, actor=user, action="sponsor.member_removed", entity_type="sponsor_member", entity_id=member.id,
          event_id=sponsor.event_id, details={"email": member.email}, request=request)
    db.delete(member)
    db.commit()
    return Response(status_code=204)


@router.get("/badge-scans")
def badge_scans(limit: int = Query(30, ge=1, le=200), sponsor: Sponsor = Depends(current_sponsor), db: Session = Depends(get_db)):
    member_ids = [m.id for m in sponsor.members]
    if not member_ids:
        return []
    scans = db.scalars(
        select(BadgeScan).where(BadgeScan.member_id.in_(member_ids)).order_by(BadgeScan.scanned_at.desc()).limit(limit)
    ).all()
    return [{"id": s.id, "member": s.member.full_name, "scanned_at": s.scanned_at} for s in scans]

"""Read attendee rows out of an uploaded .xlsx or .csv for bulk pre-registration.

Columns are matched by header name (English or Arabic, any casing/punctuation). Without a recognisable
header row, the first six columns are taken as name, email, mobile, SCFHS, National ID, profession.
"""

from __future__ import annotations

import csv
import io
import re

from openpyxl import load_workbook

from .errors import ApiError
from .validators import PROFESSIONS

MAX_ROWS = 2000
FIELDS = ("full_name", "email", "mobile", "scfhs_number", "national_id", "profession")

_HEADERS: dict[str, tuple[str, ...]] = {
    "full_name": ("fullname", "name", "attendeename", "attendee", "الاسم", "الاسمالكامل", "الاسمالثلاثي", "اسمالمشارك"),
    "email": ("email", "emailaddress", "mail", "البريد", "البريدالإلكتروني", "البريدالالكتروني", "الايميل", "الإيميل"),
    "mobile": ("mobile", "mobilenumber", "phone", "phonenumber", "tel", "الجوال", "رقمالجوال", "الهاتف", "رقمالهاتف"),
    "scfhs_number": ("scfhs", "scfhsno", "scfhsnumber", "scfhsregistrationnumber", "scfhsregistration",
                     "رقمالهيئة", "رقمالتسجيلفيالهيئة", "رقمتسجيلالهيئة"),
    "national_id": ("nationalid", "nationalidiqama", "nationalidiqamanumber", "iqama", "iqamanumber", "idnumber",
                    "id", "nid", "رقمالهوية", "الهوية", "رقمالهويةالإقامة", "رقمالهويةالاقامة", "الإقامة", "الاقامة"),
    "profession": ("profession", "job", "jobtitle", "role", "المهنة", "الوظيفة", "التخصص"),
}
_PROFESSION_LABELS = {
    "consultant": "consultant", "استشاري": "consultant", "استشارية": "consultant",
    "specialist": "specialist", "أخصائي": "specialist", "اخصائي": "specialist", "أخصائية": "specialist",
    "resident": "resident", "طبيبمقيم": "resident", "مقيم": "resident",
    "gp": "gp", "generalpractitioner": "gp", "طبيبعام": "gp",
    "nurse": "nurse", "nursing": "nurse", "تمريض": "nurse", "ممرض": "nurse", "ممرضة": "nurse",
    "pharmacist": "pharmacist", "pharmacy": "pharmacist", "صيدلي": "pharmacist", "صيدلانية": "pharmacist", "صيدلة": "pharmacist",
    "dentist": "dentist", "dental": "dentist", "طبيبأسنان": "dentist", "أسنان": "dentist",
    "alliedhealth": "allied_health", "allied_health": "allied_health", "مهنصحيةمساندة": "allied_health",
    "student": "student", "طالب": "student", "طالبة": "student",
    "other": "other", "أخرى": "other", "اخرى": "other",
}


def _norm(value: object) -> str:
    return re.sub(r"[^\w]", "", str(value or "").strip().lower())


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # Excel stores ID numbers as floats
    return str(value).strip()


def _rows_from_upload(filename: str, data: bytes) -> list[list[str]]:
    name = (filename or "").lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        text = data.decode("utf-8-sig", errors="replace")
        return [[_cell(c) for c in row] for row in csv.reader(io.StringIO(text))]
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True, max_row=MAX_ROWS + 20)]
        wb.close()
        return rows
    raise ApiError(422, "unsupported_file", "Upload an Excel (.xlsx) or CSV file.")


def profession_key(value: str) -> str | None:
    if not value:
        return None
    key = _norm(value)
    if key in PROFESSIONS:
        return key
    return _PROFESSION_LABELS.get(key)


def parse_attendees(filename: str, data: bytes) -> list[tuple[int, dict[str, str]]]:
    """Returns (spreadsheet row number, {field: raw value}) for each non-empty data row."""
    rows = [r for r in _rows_from_upload(filename, data)]
    # First row with at least two non-empty cells is the header candidate.
    header_index = next((i for i, r in enumerate(rows) if sum(1 for c in r if c) >= 2), None)
    if header_index is None:
        return []
    header = [_norm(c) for c in rows[header_index]]
    mapping: dict[int, str] = {}
    for col, name in enumerate(header):
        for field, aliases in _HEADERS.items():
            if name in aliases and field not in mapping.values():
                mapping[col] = field
    if len(mapping) >= 3:
        first_data = header_index + 1
    else:
        # No usable header: assume the standard column order and treat every row as data.
        mapping = dict(enumerate(FIELDS))
        first_data = header_index
    out = []
    for i in range(first_data, len(rows)):
        row = rows[i]
        if not any(c for c in row):
            continue
        record = {field: (row[col] if col < len(row) else "") for col, field in mapping.items()}
        out.append((i + 1, record))  # 1-based, as shown in Excel
        if len(out) >= MAX_ROWS:
            break
    return out
